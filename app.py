import os
from werkzeug.utils import secure_filename
from flask import Flask, render_template, request, jsonify, send_file, redirect, url_for, session, flash, request
from data import doc, EoA, inst, Intake, Name, NRI, site, Merge, University, Bank, Closure
import requests
from flask_pymongo import PyMongo
from docx import Document

app = Flask(__name__)
app.config["MONGO_URI"] = "mongodb+srv://Cipher:riKXPIASClOaF7sm@cluster0.iqltodr.mongodb.net/registrations"
mongodb_client = PyMongo(app)
db = mongodb_client.db
#  mongodb = riKXPIASClOaF7sm
translate_api_url = "http://349e-34-141-145-71.ngrok-free.app/"
chatbot_api_url = "http://e0a6-34-125-25-170.ngrok-free.app/"

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/translate', methods=['POST'])
def translate():  
    data = request.get_json()
    texts = data.get('texts', [])
    target_lang = data.get('target_lang', 'en')
    response = requests.post(f"{translate_api_url}/translate", json={'texts': texts, 'target_lang': target_lang})
    if response.status_code == 200:
        translated_texts = response.json().get('translated_texts', [])
        return jsonify({'translated_texts': translated_texts})
    else:
        return jsonify({'error': 'Failed to get translation'})
    
@app.route('/ask', methods=['POST'])
def ask():
    questions = request.get_json().get('question')
    response = requests.post(f"{chatbot_api_url}/input_bot", json={'questions': questions})
    if response.status_code == 200:
        answers = response.json().get('answer')
        return jsonify({'answers': answers})
    else:
        return jsonify({'error': 'Failed to get answers'})

@app.route('/register', methods=['GET', 'POST'])
def register():
    return render_template('register.html')

@app.route('/login', methods=['GET', 'POST'])
def login():
    return render_template('login.html')

@app.route('/newInsitute', methods=['GET', 'POST'])
def newInstitute():
    return render_template('newInstitute.html')

@app.route('/change', methods=['GET', 'POST'])
def change():
    return render_template('changep.html')

@app.route('/dashboard', methods=['GET', 'POST'])
def dashboard():
    return render_template('dashboard.html')

@app.route('/document-verify', methods=['GET', 'POST'])
def document_verify():
    return render_template('document-verify.html')

@app.route('/addingFile', methods=['GET', 'POST'])
def addingFile():
    return render_template('addingFile.html', doc = doc, title = 'Applications types')
@app.route('/addingFile',methods = ['GET','POST'])
def apply():
    if request.method == 'POST':
        global appType
        appType = request.form.get("application for")
        if 'pressed' in request.form:
            return redirect(url_for("pdfDisp"))
count = 0
value = list()
@app.route("/pdf", methods=['GET', 'POST'])
def pdfDisp():
    index: list = list()
    global count
    if count == 0 and not request.form.keys() and request.method == 'GET' and appType:
            # Use a dictionary to map app_type to index
            index_dict = {
                'EoA': EoA,
                'NRI': NRI,
                'site': site,
                'Merge': Merge,
                'Intake': Intake,
                'Closure': Closure,
                'inst': inst,
                'Name': Name,
                'University': University,
                'Bank': Bank
            }
            index = index_dict[appType]
            count += 1
            global value
            value = index
            app.config['MAX_CONTENT_LENGTH'] = len(value)*1024*1024
            return render_template("open.html", index=index, doc=doc, title='Upload Your Files')
    count = 0
    if request.files.keys() : 
                for i in value:
                    file = request.files[i]
                    filename = secure_filename(file.filename)
                    file.save(os.path.join(app.config['UPLOAD_FOLDER'], filename))
                    session['filename'] = filename
                return redirect(url_for('fileStore'))
    count = 0
    return redirect(url_for('pdfDisp'))

@app.route("/home", methods = ['GET','POST'])
def fileStore():
    file = session.get('file',None)
    global count 
    count = 0
    return 'Received and saved file: ' + str(file)

@app.route('/affidavit', methods=['GET', 'POST'])
def affidavit():
    return render_template('affidavit.html')
def generate_word_file(
    name,
    mobile,
    mail,
    designation,
    trust,
    parent,
    age,
    address,
    institution,
    date,
    userid,
    transactionid,
    transactiondate,
    executant,
    edesignation,
    eaddress,
    vdate,
    vmonth,
    vyear
):
    # Load the Word template
    doc = Document("AFFIDAVIT.docx")

    # Replace placeholders with user data
    for paragraph in doc.paragraphs:
        if "{{name}}" in paragraph.text:
            paragraph.text = paragraph.text.replace("{{name}}", name)

        if "{{mobile}}" in paragraph.text:
            paragraph.text = paragraph.text.replace("{{mobile}}",mobile)

        if "{{mail}}" in paragraph.text:
            paragraph.text = paragraph.text.replace("{{mail}}",mail)

        if "{{designation}}" in paragraph.text:
            paragraph.text = paragraph.text.replace("{{designation}}", designation)

        if "{{trust}}" in paragraph.text:
            paragraph.text = paragraph.text.replace("{{trust}}", trust)

        if "{{parent}}" in paragraph.text:
            paragraph.text = paragraph.text.replace("{{parent}}", parent)

        if "{{age}}" in paragraph.text:
            paragraph.text = paragraph.text.replace("{{age}}", age)

        if "{{address}}" in paragraph.text:
            paragraph.text = paragraph.text.replace("{{address}}", address)

        if "{{institution}}" in paragraph.text:
            paragraph.text = paragraph.text.replace("{{institution}}", institution)

        if "{{date}}" in paragraph.text:
            paragraph.text = paragraph.text.replace("{{date}}", date)

        if "{{userid}}" in paragraph.text:
            paragraph.text = paragraph.text.replace("{{userid}}", userid)

        if "{{transactionid}}" in paragraph.text:
            paragraph.text = paragraph.text.replace("{{transactionid}}", transactionid)

        if "{{transactiondate}}" in paragraph.text:
            paragraph.text = paragraph.text.replace("{{transactiondate}}",transactiondate)

        if "{{executant}}" in paragraph.text:
            paragraph.text = paragraph.text.replace("{{executant}}",executant)

        if "{{edesignation}}" in paragraph.text:
            paragraph.text = paragraph.text.replace("{{edesignation}}",edesignation)

        if "{{eaddress}}" in paragraph.text:
            paragraph.text = paragraph.text.replace("{{eaddress}}",eaddress)

        if "{{vdate}}" in paragraph.text:
            paragraph.text = paragraph.text.replace("{{vdate}}",vdate)

        if "{{vmonth}}" in paragraph.text:
            paragraph.text = paragraph.text.replace("{{vmonth}}",vmonth)

        if "{{vyear}}" in paragraph.text:
            paragraph.text = paragraph.text.replace("{{vyear}}",vyear)

            # Save the filled Word document
            doc.save("static/filled_template.docx")

    return "static/filled_template.docx"

@app.route('/generate-affidavit', methods=['GET', 'POST'])
def generate_affidavit():
    if request.method == "POST":
        name = request.form["name"]
        mobile = request.form["mobile"]
        mail = request.form["mail"]
        designation = request.form["designation"]
        trust = request.form["trust"]
        parent = request.form["parent"]
        age = request.form["age"]
        address = request.form["address"]
        institution = request.form["institution"]
        date = request.form["date"]
        userid = request.form["userid"]
        transactionid = request.form["transactionid"]
        transactiondate = request.form["transactiondate"]
        executant = request.form["executant"]
        edesignation = request.form["edesignation"]
        eaddress = request.form["eaddress"]
        vdate = request.form["vdate"]
        vmonth = request.form["vmonth"]
        vyear = request.form["vyear"]
        word_file_path = generate_word_file(
            name,
            mobile,
            mail,
            designation,
            trust,
            parent,
            age,
            address,
            institution,
            date,
            userid,
            transactionid,
            transactiondate,
            executant,
            edesignation,
            eaddress,
            vdate,
            vmonth,
            vyear,
        )
        # Provide the generated Word file for download
        return send_file(word_file_path, as_attachment=True)

    return render_template("generate_affidavit.html")


@app.route('/login-data', methods=['POST'])
def login_data():
    try:
        email = request.json.get('username')
        password = request.json.get('password')
        user = db.institutes.find_one({"email": email, "password": password})
        if user:
            return jsonify({'message': 'Success'}), 200
        else:
            return jsonify({'message': 'Failed'}), 401
        
    except Exception as e:
        print("Error:", e)
        return jsonify({'error': str(e)})
    
@app.route('/register-data', methods=['POST'])
def register_data():
    try:
        if request.method == 'POST':
            collegeName = request.json.get('collegeName')
            instituteType = request.json.get('instituteType')
            state = request.json.get('state')
            email = request.json.get('email')
            contactNumber = request.json.get('contactNumber')
            collegeAddress = request.json.get('collegeAddress')
            postalCode = request.json.get('postalCode')
            password = request.json.get('password')
            checkbox = request.json.get('checkbox', False)

            db.institutes.insert_one({
                'collegeName': collegeName,
                'instituteType': instituteType,
                'state': state,
                'email': email,
                'contactNumber': contactNumber,
                'collegeAddress': collegeAddress,
                'postalCode': postalCode,
                'password': password,
                'checkbox': checkbox,
            })

            return jsonify({'message': 'success'}), 200
        else:
            return jsonify({'error': 'Invalid request method'}), 401
    except Exception as e:
        print("Error:", e)
        return jsonify({'error': str(e)})



if __name__ == '__main__':
    app.run(debug=True)
